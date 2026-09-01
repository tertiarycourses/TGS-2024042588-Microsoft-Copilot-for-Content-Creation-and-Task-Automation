#!/usr/bin/env python3
"""Generate the Learner Guide as BOTH a Markdown mirror (LEARNER-GUIDE.md at the
repo root) and a DOCX (courseware/LG-*.docx) from one content stream, so the two
never diverge.

The LG is where the DETAILED STEP-BY-STEP lives — the slide deck deliberately
carries none.
"""
import os, sys, glob
HERE=os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0,HERE)
REPO=os.path.dirname(HERE)
import course_data as C
from data_domain1 import DOMAIN1; from data_domain2 import DOMAIN2; from data_domain3 import DOMAIN3
from data_domain4 import DOMAIN4; from data_domain5 import DOMAIN5
import prodoc
from docx import Document
from docx.shared import Pt, RGBColor, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
ACT=DOMAIN1+DOMAIN2+DOMAIN3+DOMAIN4+DOMAIN5
ASSETS=os.path.join(REPO,"courseware","assets")

# ---------------- block DSL (one content stream -> MD + DOCX) ----------------
B=[]
def h1(t): B.append(("h1",t))
def h2(t): B.append(("h2",t))
def h3(t): B.append(("h3",t))
def p(t):  B.append(("p",t))
def bullets(xs): B.append(("bullets",list(xs)))
def steps(xs): B.append(("steps",list(xs)))
def quote(t): B.append(("quote",t))
def note(t): B.append(("note",t))
def table(headers,rows): B.append(("table",headers,rows))
def rule(): B.append(("rule",))
def figure(path,caption): B.append(("figure",path,caption))

def lab_folder(n):
    g=glob.glob(os.path.join(REPO,"labs",f"lab-{n:02d}-*"))
    return os.path.basename(g[0]) if g else f"lab-{n:02d}"

# ================================================================= CONTENT
h1("How to Use This Guide")
p(f"This Learner Guide accompanies the WSQ course {C.TITLE} ({C.COURSE_CODE}), conducted by "
  f"{C.ORG}. It contains the full step-by-step procedure for all {len(ACT)} hands-on labs. The "
  "slide deck used in class is deliberately visual and concept-led — every detailed procedure "
  "lives here and in the instruction PDF inside each lab folder.")
p("All twenty-five labs follow one connected scenario: Lumina Living Pte Ltd, a Singapore "
  "workspace-furniture company, is launching its Sustainable Workspace campaign. Days 1 and 2 "
  "create and verify the campaign's content with Microsoft 365 Copilot. Days 3 to 5 automate the "
  "process behind it with Power Automate and Copilot Studio agents, then integrate, test, secure "
  "and govern the result.")
h2("Course and lab file access")
bullets([
 "Course materials are on the LMS/TMS portal at https://lms-tms.tertiaryinfotech.com.",
 "Each lab has its own folder under labs/ containing an Excel workbook, its mock data and a printable instruction PDF.",
 "You may use this guide, the slides and your completed lab files during the open-book assessment.",
])
h2("How each lab section is organised")
bullets([
 "Before you begin — what must already be in place.",
 "Primary prompt — the exact prompt to use, which you may adapt.",
 "Detailed procedure — the numbered steps, with the reason for each.",
 "Acceptance test — how you prove the lab is complete.",
 "Evidence to retain — what to keep for the practical assessment.",
])

h1("Course Learning Outcomes")
bullets(C.LEARNING_OUTCOMES)

h1("Assessed Knowledge and Abilities")
p("The Written Assessment covers every knowledge statement (K); the Practical Performance covers "
  "every ability statement (A). Each lab states the codes it develops.")
table(["Code","Statement"],[[k,v] for k,v in list(C.KNOWLEDGE)+list(C.ABILITIES)])

h1("The Prompt and Verification Framework")
p("Every prompt in this course follows a four-part contract, and every output is verified before "
  "it is used. This framework is assessed, so apply it in every lab.")
h3("The four-part prompt contract")
table(["Part","What it states","Example"],
 [["Goal","The outcome you want produced","Draft a two-page project brief"],
  ["Context","Who it is for and why","For the leadership team, to approve a budget"],
  ["Source","What the answer must come from","Using the attached workbook and template"],
  ["Expectation","Format, length, tone and limits","Three measurable objectives; mark unevidenced points as open questions"]])
h3("The verification duty")
p("Copilot proposes; you dispose. Before any output is released, every claim must be traced to a "
  "source, and a named person must accept it.")
bullets([
 "Trace each quantitative claim to a named sheet and cell range, or to a document section.",
 "Rewrite any claim you cannot trace as an open question — never release it as a fact.",
 "Check what was omitted, not only what was stated: a true summary can still mislead.",
 "Record the reviewer and the decision. Accountability never transfers to the model.",
])
h3("Three ways a generated draft goes wrong")
table(["Failure","What it looks like","How you catch it"],
 [["Invention","A confident fact that appears in no source","Trace every claim to its source"],
  ["Omission","A true summary that drops the caveat","Read the source, not only the summary"],
  ["Misframing","Accurate numbers arranged to imply the wrong thing","Ask what decision a reader would make"]])

h1("Workbook Conventions Used in Every Lab")
p("Every lab workbook carries the same control sheets so evidence is captured consistently and "
  "can be assessed the same way across all twenty-five labs.")
table(["Sheet","Purpose"],
 [["Scenario","The lab's goal, alignment, apps and acceptance criterion"],
  ["Steps","The numbered procedure, with a Done column to track progress"],
  ["Prompt_Log","Every prompt used, its purpose, its grounding source and its reviewer"],
  ["Evidence_Register","Each artifact, its source, the check applied and the result"],
  ["Acceptance_Check","The acceptance criterion, the evidence for it and the reviewer"]])
p("Labs that work with data add their own sheets — for example Campaign_Data, Assumptions and "
  "Controls for the Excel labs, and test-result sheets for the agent labs.")

# ---------------- per-topic, per-lab ----------------
for t in C.TOPICS:
    labs=[a for a in ACT if a["topic"]==t["num"]]
    h1(f"Topic {t['code']} — {t['title']}")
    p(f"{t['subtitle']}  ·  {t['days']}")
    h3("Key concepts")
    bullets(t["concepts"])
    for a in labs:
        h2(f"Lab {a['num']:02d} — {a['title']}")
        p(f"Alignment: {a['km']} · {a['am']}.  Level: {a['level']}.  Apps: {a['apps']}.")
        p(f"Objective: {a['objective']}.")
        p(f"Goal: {a['desc']}")
        h3("What you will build")
        p(a["build"])
        fig=os.path.join(ASSETS,"figures",f"lab-{a['num']:02d}-workflow.png")
        if os.path.exists(fig):
            figure(fig, f"Figure {a['num']}. Lab {a['num']:02d} workflow — the four stages and the artifact produced.")
        h3("Before you begin")
        bl=[f"Open the lab folder labs/{lab_folder(a['num'])}.",
            "Open the lab workbook and read the Scenario sheet.",
            "Have the printable instruction PDF open or to hand."]
        for fn,d in a["files"]:
            bl.append(f"Confirm the file {fn} is present — {d}.")
        bullets(bl)
        h3("Primary prompt")
        quote(a["prompt"])
        h3("Detailed procedure")
        steps(a["steps"])
        h3("Acceptance test")
        p(a["test"])
        h3("Evidence to retain")
        bullets([a["evidence"],
                 "Your completed workbook, including the Prompt_Log and Evidence_Register sheets.",
                 "A credential-safe screenshot — never capture passwords or personal data."])
        rule()

h1("Quick Prompt Reference")
p("Every primary prompt used in the course, in lab order.")
table(["Lab","Prompt"],[[f"{a['num']:02d}",a["prompt"]] for a in ACT])

h1("Assessment")
p("The course is assessed by two instruments, both open book. You may use this guide, the slide "
  "deck and your completed lab files.")
bullets([C.ASSESSMENT["written"],C.ASSESSMENT["practical"],C.ASSESSMENT["note"],
 "Remember to take the Assessment digital attendance (TRAQOM) and complete the TRAQOM survey.",
 "Submit your answers on the LMS at https://lms-tms.tertiaryinfotech.com.",
 "A re-assessment and appeal process is available if required."])
h3("On assessment day")
steps([("Complete the TRAQOM survey","Scan the QR code shown on the LMS."),
       ("Take the assessment digital attendance","Scan the SSG QR code the trainer displays."),
       ("Sit the Written Assessment (SAQ)","One hour, open book, individual work."),
       ("Sit the Practical Performance (PP)","One hour, open book, individual work."),
       ("Submit your answers on the LMS","Check that every question has been answered."),
       ("Sign the Assessment Summary Record","Confirm your result has been recorded.")])

h1("Support")
bullets([f"Course portal: https://lms-tms.tertiaryinfotech.com",
         f"Course page: {C.COURSE_URL}",
         f"Provider: {C.ORG} ({C.UEN})"])

# ================================================================= RENDERERS
def render_md(path):
    L=[f"# {C.TITLE} — Learner Guide\n",
       f"**Course code:** {C.COURSE_CODE}  ·  **Version:** {C.VERSION}  ·  {C.VERSION_DATE}  \n"
       f"**Conducted by:** {C.ORG} ({C.UEN})\n"]
    for b in B:
        k=b[0]
        if   k=="h1": L.append(f"\n## {b[1]}\n")
        elif k=="h2": L.append(f"\n### {b[1]}\n")
        elif k=="h3": L.append(f"\n**{b[1]}**\n")
        elif k=="p":  L.append(b[1]+"\n")
        elif k=="quote": L.append("> "+b[1]+"\n")
        elif k=="note":  L.append(f"> **Note:** {b[1]}\n")
        elif k=="bullets":
            L.extend(f"- {x}" for x in b[1]); L.append("")
        elif k=="steps":
            for i,(ttl,d) in enumerate(b[1],1):
                L.append(f"{i}. **{ttl}**  "); L.append(f"   {d}")
            L.append("")
        elif k=="table":
            hs,rows=b[1],b[2]
            L.append("| "+" | ".join(hs)+" |")
            L.append("|"+"|".join(["---"]*len(hs))+"|")
            for r in rows: L.append("| "+" | ".join(str(c).replace("|","\\|") for c in r)+" |")
            L.append("")
        elif k=="figure":
            rel=os.path.relpath(b[1], REPO)
            L.append(f"![{b[2]}]({rel})\n")
            L.append(f"*{b[2]}*\n")
        elif k=="rule": L.append("\n---\n")
    L.append(f"\n---\n\n*{C.TITLE} · {C.COURSE_CODE} · Version {C.VERSION} · © 2026 {C.ORG}*")
    open(path,"w").write("\n".join(L))

def render_docx(path):
    doc=Document()
    st=doc.styles["Normal"]; st.font.name="Arial"; st.font.size=Pt(11)
    for s in doc.sections:
        s.top_margin=Inches(0.9); s.bottom_margin=Inches(0.9)
        s.left_margin=Inches(1.0); s.right_margin=Inches(1.0)
    prodoc.style_headings(doc)
    prodoc.add_cover_page(doc,"Learner Guide",C.TITLE,C.VERSION.lstrip("v"),
        org_logo=os.path.join(ASSETS,"tertiary-infotech-logo.png"),
        course_code=C.COURSE_CODE)
    prodoc.add_version_control(doc,[
        (C.VERSION.lstrip("v"), C.VERSION_DATE,
         f"First release of the five-day Learner Guide covering all {len(ACT)} labs — Copilot for "
         "content creation (Days 1-2) and Copilot for workflow and agents (Days 3-5), aligned to "
         "the v1.0 slide deck, Lesson Plan and the WA/PP assessment.", C.TRAINER)])
    prodoc.add_toc(doc)
    for b in B:
        k=b[0]
        if   k=="h1": doc.add_heading(b[1],level=1)
        elif k=="h2": doc.add_heading(b[1],level=2)
        elif k=="h3": doc.add_heading(b[1],level=3)
        elif k=="p":  doc.add_paragraph(b[1])
        elif k=="quote":
            pr=doc.add_paragraph(); r=pr.add_run(b[1]); r.italic=True
            pr.paragraph_format.left_indent=Inches(0.35)
        elif k=="note":
            pr=doc.add_paragraph(); r=pr.add_run("Note: "+b[1])
            r.italic=True; r.font.color.rgb=prodoc.GREY
        elif k=="bullets":
            for x in b[1]: doc.add_paragraph(str(x),style="List Bullet")
        elif k=="steps":
            for i,(ttl,d) in enumerate(b[1],1):
                pr=doc.add_paragraph(); r=pr.add_run(f"Step {i}. {ttl}")
                r.bold=True; pr.paragraph_format.space_before=Pt(8)
                pr.paragraph_format.space_after=Pt(2)
                doc.add_paragraph(d)
        elif k=="table":
            hs,rows=b[1],b[2]
            t=doc.add_table(rows=0,cols=len(hs)); t.style="Table Grid"
            hr=t.add_row().cells
            for i,h in enumerate(hs):
                hr[i].text=""; r=hr[i].paragraphs[0].add_run(h)
                r.bold=True; r.font.size=Pt(9.5); r.font.color.rgb=RGBColor(0xFF,0xFF,0xFF)
                prodoc._shade_cell(hr[i],"1F6FEB")
            for row in rows:
                c=t.add_row().cells
                for i,v in enumerate(row):
                    c[i].text=""; rr=c[i].paragraphs[0].add_run(str(v)); rr.font.size=Pt(9.5)
            doc.add_paragraph()
        elif k=="figure":
            ip=doc.add_paragraph(); ip.alignment=WD_ALIGN_PARAGRAPH.CENTER
            ip.add_run().add_picture(b[1], width=Inches(6.3))
            cp=doc.add_paragraph(); cp.alignment=WD_ALIGN_PARAGRAPH.CENTER
            cr=cp.add_run(b[2]); cr.font.size=Pt(9); cr.italic=True
            cr.font.color.rgb=prodoc.GREY
        elif k=="rule":
            pr=doc.add_paragraph(); pr.paragraph_format.space_after=Pt(2)
            r=pr.add_run("_"*88); r.font.color.rgb=RGBColor(0xDD,0xE3,0xEC); r.font.size=Pt(8)
    prodoc.add_page_numbers(doc); prodoc.enable_update_fields(doc)
    doc.save(path)

if __name__=="__main__":
    md=os.path.join(REPO,"LEARNER-GUIDE.md")
    dx=os.path.join(REPO,"courseware",f"LG-{C.SHORT_TITLE}.docx")
    render_md(md); render_docx(dx)
    print("Saved",md); print("Saved",dx)
    print(f"blocks {len(B)}, labs {len(ACT)}")
