#!/usr/bin/env python3
"""Generate one self-contained folder per lab: workbook + mock data + README + instruction PDF."""
import os, sys, re, subprocess, shutil
HERE=os.path.dirname(os.path.abspath(__file__)); sys.path.insert(0,HERE)
REPO=os.path.dirname(HERE)
import course_data as C
from data_domain1 import DOMAIN1; from data_domain2 import DOMAIN2; from data_domain3 import DOMAIN3
from data_domain4 import DOMAIN4; from data_domain5 import DOMAIN5
import mockdata as M
ACT=DOMAIN1+DOMAIN2+DOMAIN3+DOMAIN4+DOMAIN5

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.utils import get_column_letter
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

LABS=os.path.join(REPO,"labs")
HDR_FILL=PatternFill("solid",fgColor="1F6FEB"); HDR_FONT=Font(color="FFFFFF",bold=True,name="Arial",size=10)
TITLE_FONT=Font(bold=True,size=14,name="Arial",color="161B26")
BODY=Font(name="Arial",size=10)
THIN=Side(style="thin",color="D6DEE8"); BOX=Border(left=THIN,right=THIN,top=THIN,bottom=THIN)

def slug(t):
    s=re.sub(r"[^a-z0-9]+","-",t.lower()).strip("-")
    return re.sub(r"-+","-",s)

def _hdr(ws,headers,row=1):
    for i,h in enumerate(headers,1):
        c=ws.cell(row=row,column=i,value=h); c.fill=HDR_FILL; c.font=HDR_FONT
        c.alignment=Alignment(horizontal="left",vertical="center"); c.border=BOX
    ws.freeze_panes=ws.cell(row=row+1,column=1)
    ws._current_row=row          # so a later ws.append() lands on row+1, not row+2

def _widths(ws,widths):
    for i,w in enumerate(widths,1): ws.column_dimensions[get_column_letter(i)].width=w

def _table(ws,name,nrows,ncols,row0=1):
    ref=f"A{row0}:{get_column_letter(ncols)}{row0+nrows}"
    t=Table(displayName=name,ref=ref)
    t.tableStyleInfo=TableStyleInfo(name="TableStyleMedium2",showRowStripes=True)
    ws.add_table(t)

def scenario_sheet(wb,a):
    ws=wb.active; ws.title="Scenario"
    ws["A1"]=f"Lab {a['num']:02d}: {a['title']}"; ws["A1"].font=TITLE_FONT
    ws["A2"]="Lumina Living Pte Ltd · Sustainable Workspace campaign"; ws["A2"].font=Font(name="Arial",size=10,italic=True,color="5B6372")
    rows=[("Course",C.TITLE),("Course code",C.COURSE_CODE),("Topic",f"{a['topic']:02d}"),
          ("Day",str(a['day'])),("Level",a['level']),("Alignment",f"{a['km']} · {a['am']}"),
          ("Primary apps",a['apps']),("Goal",a['desc']),("You will build",a['build']),
          ("Acceptance test",a['test'])]
    r=4
    for k,v in rows:
        ws.cell(row=r,column=1,value=k).font=Font(name="Arial",size=10,bold=True,color="1F6FEB")
        c=ws.cell(row=r,column=2,value=v); c.font=BODY; c.alignment=Alignment(wrap_text=True,vertical="top")
        r+=1
    _widths(ws,[20,104]); return ws

def common_sheets(wb,a):
    ws=wb.create_sheet("Prompt_Log")
    _hdr(ws,["Prompt_ID","Purpose","Prompt","Grounding","Reviewer","Status"])
    ws.append([f"P{a['num']:02d}-01",a['title'],a['prompt'],"Attached lab files only","Learner","Not started"])
    for i in range(2,5): ws.append([f"P{a['num']:02d}-{i:02d}","","","Attached lab files only","Learner","Not started"])
    _widths(ws,[12,34,80,26,14,14]); _table(ws,f"PromptLog{a['num']:02d}",4,6)
    for row in ws.iter_rows(min_row=2,max_row=5,max_col=6):
        for c in row: c.font=BODY; c.alignment=Alignment(wrap_text=True,vertical="top")

    ws=wb.create_sheet("Evidence_Register")
    _hdr(ws,["Evidence_ID","Artifact","Source","Check","Result","Owner","Notes"])
    checks=["Source and scope verified","Quantitative claims checked against source",
            "Sensitive or personal data reviewed","Output approved by a named person"]
    for i,ck in enumerate(checks,1): ws.append([f"E{i:02d}","","",ck,"","Learner",""])
    _widths(ws,[13,30,30,42,14,14,30]); _table(ws,f"Evidence{a['num']:02d}",len(checks),7)
    for row in ws.iter_rows(min_row=2,max_row=1+len(checks),max_col=7):
        for c in row: c.font=BODY; c.alignment=Alignment(wrap_text=True,vertical="top")

    ws=wb.create_sheet("Acceptance_Check")
    _hdr(ws,["Criterion","Evidence","Status","Reviewer"])
    ws.append([a['test'],"See Evidence_Register","Not tested","Learner"])
    _widths(ws,[90,26,16,16]); _table(ws,f"Accept{a['num']:02d}",1,4)
    for row in ws.iter_rows(min_row=2,max_row=2,max_col=4):
        for c in row: c.font=BODY; c.alignment=Alignment(wrap_text=True,vertical="top")

    ws=wb.create_sheet("Steps")
    _hdr(ws,["#","Step","What to do","Done"])
    for i,(t,d) in enumerate(a['steps'],1): ws.append([i,t,d,""])
    _widths(ws,[5,46,86,9]); _table(ws,f"Steps{a['num']:02d}",len(a['steps']),4)
    for row in ws.iter_rows(min_row=2,max_row=1+len(a['steps']),max_col=4):
        for c in row: c.font=BODY; c.alignment=Alignment(wrap_text=True,vertical="top")

def add_campaign(wb, with_derived=True):
    ws=wb.create_sheet("Campaign_Data")
    hdrs=list(M.CAMPAIGN_HEADERS)
    if with_derived: hdrs+=["Margin","Margin_Percent","Conversion_Rate","Return_Rate","Target_Status"]
    _hdr(ws,hdrs)
    rows=M.campaign_rows()
    for r in rows: ws.append(r+([""]*(len(hdrs)-len(r))))
    _widths(ws,[12,8,10,10,16]+[13]*(len(hdrs)-5)); _table(ws,"CampaignData",len(rows),len(hdrs))
    for row in ws.iter_rows(min_row=2,max_row=1+len(rows),max_col=len(hdrs)):
        for c in row: c.font=BODY
    ws2=wb.create_sheet("Assumptions")
    _hdr(ws2,["Assumption","Value","Purpose"])
    for r in [["Target_Margin",0.32,"On-target margin threshold"],
              ["Target_Conversion",0.04,"Conversion threshold"],
              ["High_Return_Rate",0.08,"Return rate review threshold"],
              ["Monthly_Growth",0.045,"Base-case month-on-month growth"],
              ["Downside_Factor",0.82,"Downside scenario multiplier"],
              ["Upside_Factor",1.18,"Upside scenario multiplier"]]: ws2.append(r)
    _widths(ws2,[24,14,44]); _table(ws2,"Assump",6,3)
    for row in ws2.iter_rows(min_row=2,max_row=7,max_col=3):
        for c in row: c.font=BODY
    ws3=wb.create_sheet("Controls")
    _hdr(ws3,["Control","Result","Expected","Status"])
    for r in [["Record count","",96,""],["Blank keys","",0,""],
              ["Negative revenue","",0,""],["Error values in derived columns","",0,""]]: ws3.append(r)
    _widths(ws3,[36,16,14,14]); _table(ws3,"Controls",4,4)
    for row in ws3.iter_rows(min_row=2,max_row=5,max_col=4):
        for c in row: c.font=BODY

def extra_sheet(wb,name,headers,rows,widths,tname):
    ws=wb.create_sheet(name); _hdr(ws,headers)
    for r in rows: ws.append(r)
    _widths(ws,widths); _table(ws,tname,max(len(rows),1),len(headers))
    for row in ws.iter_rows(min_row=2,max_row=1+max(len(rows),1),max_col=len(headers)):
        for c in row: c.font=BODY; c.alignment=Alignment(wrap_text=True,vertical="top")
    return ws

# ---- per-lab extra data sheets -------------------------------------------------
def lab_specific(wb,a):
    n=a['num']
    if n==1:
        extra_sheet(wb,"Setup",["Item","Value","Verified"],
            [["Copilot surface (Chat / Work)","",""],["Licence confirmed","",""],
             ["Grounding source used","",""],["Tenant / account","",""]],[34,34,14],"Setup01")
    if n in (2,4,5): add_campaign(wb, with_derived=False)
    if n==2:
        extra_sheet(wb,"Brief_Data",["Field","Value","Source"],
            [["Campaign name","Sustainable Workspace","Marketing plan"],
             ["Launch date","2026-10-05","Marketing plan"],
             ["Budget (SGD)",185000,"Finance approval"],
             ["Target revenue (SGD)",940000,"Finance approval"],
             ["Primary segment","Corporate offices, 50-250 staff","Segment study"],
             ["Milestone 1","Creative sign-off 2026-09-12","Project plan"],
             ["Milestone 2","Channel launch 2026-10-05","Project plan"],
             ["Milestone 3","Post-launch review 2026-11-16","Project plan"],
             ["Risk 1","Supplier lead time on Acoustic Pods","Risk log"],
             ["Risk 2","Retail staff training not scheduled","Risk log"]],[26,44,26],"BriefData")
    if n==3:
        extra_sheet(wb,"Issue_Log",["Issue_ID","Clause","Problem","Suggested_Rewrite","Decision","Rationale"],
            [[f"I{i:02d}","","","","",""] for i in range(1,9)],[10,20,40,44,14,30],"IssueLog")
    if n==4:
        extra_sheet(wb,"Evidence_Map",["Slide","Claim","Source_Sheet","Cell_Range","Verified"],
            [[i,"","","",""] for i in range(1,9)],[8,44,20,18,12],"EvidenceMap")
    if n==5:
        extra_sheet(wb,"Chart_Plan",["Slide","Message","Chart_Type","Axis_Units","Source_Range","Speaker_Note_Verified"],
            [[i,"","","","",""] for i in range(1,7)],[8,40,18,18,20,20],"ChartPlan")
    if n in (6,7,8): add_campaign(wb, with_derived=(n==6))
    if n==6:
        extra_sheet(wb,"Formula_Audit",["Column","Formula","What_It_Computes","Edge_Case_Checked","Reviewer"],
            [[c,"","","",""] for c in ["Margin","Margin_Percent","Conversion_Rate","Return_Rate","Target_Status"]],
            [20,40,36,24,14],"FormulaAudit")
    if n==7:
        extra_sheet(wb,"KPI_Summary",["KPI","Value","Source_Range","Comment"],
            [[k,"","",""] for k in ["Total revenue","Total cost","Overall margin %","Overall conversion %",
             "Revenue vs target","Highest region","Lowest channel"]],[26,16,20,44],"KPISummary")
        extra_sheet(wb,"Anomaly_Log",["Record_ID","Field","Observed","Why_Flagged","Action"],
            [["","","","",""] for _ in range(6)],[14,18,18,44,28],"AnomalyLog")
        extra_sheet(wb,"Findings",["Finding","Evidence_Range","Recommended_Action","Owner"],
            [["","","",""] for _ in range(5)],[54,20,44,16],"Findings")
    if n==8:
        extra_sheet(wb,"Scenarios",["Month","Base","Downside","Upside","Assumption_Ref"],
            [[m,"","","",""] for m in M.MONTHS[:6]],[10,16,16,16,20],"Scenarios")
        extra_sheet(wb,"Sensitivity",["Driver","-20%","-10%","Base","+10%","+20%"],
            [["Conversion rate","","","","",""],["Average order value","","","","",""]],
            [24,12,12,12,12,12],"Sensitivity")
        extra_sheet(wb,"Guardrails",["Condition","Threshold","Action_If_Breached","Owner"],
            [["Margin below target","","",""],["Conversion below target","","",""],
             ["Return rate above threshold","","",""],["Forecast assumption invalid","","",""]],
            [34,16,40,16],"Guardrails")
    if n==9:
        extra_sheet(wb,"Triage_Register",["Enquiry_ID","From","Subject_Summary","Classification","Owner","Due"],
            [[e[0],e[1],e[4][:60],"","",""] for e in M.ENQUIRIES],[13,20,58,18,14,12],"Triage")
        extra_sheet(wb,"Thread_Summary",["Item","Type","Stated_By","Source_Line_Verified"],
            [["","Decision","",""],["","Decision","",""],["","Open question","",""],
             ["","Open question","",""],["","Contradiction","",""]],[50,18,20,22],"ThreadSum")
    if n==10:
        extra_sheet(wb,"Decision_Log",["Decision","Supporting_Line","Made_By","Date"],
            [["","","",""] for _ in range(5)],[50,44,18,12],"DecisionLog")
        extra_sheet(wb,"Action_Register",["Action_ID","Action","Owner","Due_Date","Supporting_Line","Status"],
            [[f"A{i:02d}","","","","","Not started"] for i in range(1,7)],[10,44,18,14,40,14],"ActionReg")
    if n==11:
        extra_sheet(wb,"Environment",["Item","Value","Verified"],
            [["Environment name","",""],["Environment type","",""],["Region","",""],
             ["Dataverse database present","",""],["Date verified","",""]],[30,36,14],"EnvSheet")
        extra_sheet(wb,"Connections",["Connector","Required_For","Connected","Verified_By"],
            [["Microsoft Forms","Labs 12-15","",""],["Office 365 Outlook","Labs 12-15","",""],
             ["Excel Online (Business)","Lab 13","",""],["Approvals","Labs 15, 23","",""],
             ["SharePoint","Labs 18, 21","",""],["HTTP","Lab 21","",""]],[26,22,14,16],"Conns")
    if n in (12,13,14,15):
        extra_sheet(wb,"Flow_Design",["Element","Configuration","Dynamic_Value_Used","Verified"],
            [["Trigger","","",""],["Action 1","","",""],["Action 2","","",""],
             ["Action 3","","",""],["Condition / branch","","",""]],[22,44,30,12],"FlowDesign")
        extra_sheet(wb,"Run_Verification",["Run","Input_Used","Expected_Result","Actual_Result","Run_Status"],
            [[f"Run {i}","","","",""] for i in range(1,4)],[10,34,34,34,14],"RunVerify")
    if n==13:
        extra_sheet(wb,"EnquiryLog_Sample",M.ENQUIRY_HEADERS,M.ENQUIRIES,[13,20,32,18,58],"EnqSample")
    if n==14:
        extra_sheet(wb,"Branch_Tests",["Test","Input_Message","Expected_Branch","Actual_Branch","Pass"],
            [["Urgent keyword present","","If yes","",""],["No keyword","","If no","",""]],
            [24,44,18,18,10],"BranchTests")
    if n==15:
        extra_sheet(wb,"Approval_Tests",["Test","Applicant","Outcome_Chosen","Expected_Email","Actual","Pass"],
            [["Approved path","","Approve","",""," "],["Rejected path","","Reject","",""," "]],
            [18,20,18,34,26,8],"ApprovalTests")
    if n==16:
        extra_sheet(wb,"Design_Canvas",["Field","Value"],
            [["The one job this agent does",""],["Who uses it",""],["In scope",""],
             ["Explicitly out of scope",""],["Tone and style",""],["Escalation contact",""]],[34,84],"Canvas")
        extra_sheet(wb,"Knowledge_Register",["Source","Owner","Approved","Why_Needed"],
            [["Procurement policy","","",""],["Approved vendor list","","",""],["",""," ",""]],[34,20,12,44],"KnowReg")
        extra_sheet(wb,"Risk_Controls",["Risk","Likelihood","Impact","Control","Control_Type"],
            [["Agent answers outside scope","","","",""],["Agent invents a price","","","",""],
             ["Agent recommends a suspended vendor","","","",""],["Consequential action taken without a person","","","",""]],
            [44,14,12,44,20],"RiskCtrl")
        extra_sheet(wb,"Acceptance_Tests",["Test_ID","Question_Or_Input","Expected_Behaviour","Actual","Pass"],
            [[f"T{i:02d}","","",""," "] for i in range(1,11)],[10,44,44,34,8],"AcceptTests")
    if n in (17,19,22):
        extra_sheet(wb,"Test_Results",["Test_ID","Input","Expected","Actual","Citation_Given","Pass"],
            [[f"T{i:02d}","","","","",""] for i in range(1,11)],[10,40,36,40,26,8],"TestResults")
    if n in (16,17,18):
        extra_sheet(wb,"Vendors",M.VENDOR_HEADERS,M.VENDORS,[12,32,18,16,18],"VendorList")
    if n==18:
        extra_sheet(wb,"Tool_Contract",["Element","Value"],
            [["Tool name",""],["When the agent should call it",""],["Input field(s)",""],
             ["Output field(s)",""],["Published",""]],[34,84],"ToolContract")
        extra_sheet(wb,"Requisition_Tests",M.REQUISITION_HEADERS+["Actual_Outcome","Pass"],
            [r+["",""] for r in M.REQUISITIONS],[10,30,14,18,46,34,8],"ReqTests")
    if n==19:
        extra_sheet(wb,"Routing_Test",["Request","Expected_Specialist","Actual_Specialist","Attributed","Pass"],
            [["","","","",""] for _ in range(5)],[46,26,26,14,8],"RoutingTest")
    if n==20:
        extra_sheet(wb,"Release_Log",["Item","Value"],
            [["Agent name",""],["Version published",""],["Date",""],["Tests passed",""],
             ["Approver",""],["Channel",""],["Rollback plan",""],["Rollback decided by",""]],[28,90],"ReleaseLog")
    if n==21:
        extra_sheet(wb,"Endpoint_Contract",["Element","Value"],
            [["Endpoint URL (do not share)",""],["Request method","POST"],
             ["Request fields",""],["Response status code",""],["Response fields",""]],[32,86],"Endpoint")
        extra_sheet(wb,"Test_Calls",["Call","Request_Body","Expected_Response","Actual_Response","Pass"],
            [["Valid request","","","",""],["Missing required field","","","",""],
             ["Unknown value","","","",""]],[22,40,34,34,8],"TestCalls")
    if n==22:
        extra_sheet(wb,"RAG_Tests",["Test_ID","Question","Expected_Behaviour","Actual_Answer","Citation","Pass"],
            [[f"R{i:02d}","","","","",""] for i in range(1,11)],[10,40,32,40,26,8],"RagTests")
    if n==23:
        extra_sheet(wb,"Agency_Boundary",["Action","Agent_Alone","Needs_Human_Gate","Never_Allowed","Control"],
            [["Answer a policy question","","","",""],["Look up a vendor status","","","",""],
             ["Raise a requisition","","","",""],["Release a payment","","","",""]],[34,14,18,16,36],"Boundary")
        extra_sheet(wb,"Approval_Evidence",["Test","Request","Approver","Decision","Audit_Row_Written","Pass"],
            [["Approved","","","","",""],["Rejected","","","","",""],["Edge case","","","","",""]],
            [16,36,20,14,20,8],"ApprovalEvidence")
    if n==24:
        extra_sheet(wb,"Test_Log",["Artifact","Test","Expected","Actual","Run_ID","Pass"],
            [[f"Lab {i}","Happy path","","","",""] for i in [12,13,14,15,17,18,21,23]],[14,24,30,30,20,8],"TestLog")
        extra_sheet(wb,"Issue_Register",["Issue_ID","Artifact","Symptom","Cause_Type","Root_Cause","Fix_Applied","Retest_Result"],
            [[f"IS{i:02d}","","","","","",""] for i in range(1,7)],[10,14,34,18,34,34,18],"IssueRegister")
    if n==25:
        extra_sheet(wb,"Pipeline_Stages",["Stage","System","Owner","Output","Verified_By","Run_ID"],
            [["Request received","","","","",""],["Content generated","","","","",""],
             ["Content verified","","","","",""],["Automation run","","","","",""],
             ["Human approval","","","","",""],["Output released","","","","",""]],[22,22,16,30,18,18],"Stages")
        extra_sheet(wb,"Governance_Pack",["Evidence_Item","Location","Present","Owner"],
            [["Prompt log","","",""],["Evidence register","","",""],["Run identifiers","","",""],
             ["Approval record","","",""],["Final output","","",""],["Residual risk noted","","",""]],
            [30,34,12,18],"GovPack")

def build_workbook(a,path):
    wb=Workbook(); scenario_sheet(wb,a); common_sheets(wb,a); lab_specific(wb,a)
    wb.save(path)

# ---- README ---------------------------------------------------------------
def build_readme(a,folder):
    L=[]
    L.append(f"# Lab {a['num']:02d}: {a['title']}\n")
    L.append(f"**Topic {a['topic']:02d}** · **Day {a['day']}** · **{a['level']}** · **{a['apps']}** · **{a['km']} · {a['am']}**\n")
    L.append("## Goal\n"); L.append(a['desc']+"\n")
    L.append("## Expected output\n"); L.append(a['build']+"\n")
    L.append(f"![Lab {a['num']:02d} workflow](LAB-{a['num']:02d}-Workflow.png)\n")
    L.append(f"*Figure. Lab {a['num']:02d} workflow — the four stages and the artifact produced.*\n")
    L.append("## Files\n")
    for fn,d in a['files']: L.append(f"- `{fn}` — {d}")
    L.append(f"- `LAB-{a['num']:02d}-Instructions.pdf` — printable detailed guide")
    L.append(f"- `LAB-{a['num']:02d}-Workflow.png` — the workflow diagram above\n")
    L.append("## Primary prompt\n"); L.append("> "+a['prompt']+"\n")
    L.append("## Procedure\n")
    for i,(t,d) in enumerate(a['steps'],1):
        L.append(f"{i}. **{t}**\n"); L.append(f"   {d}\n")
    L.append("## Acceptance test\n"); L.append(a['test']+"\n")
    L.append("## Evidence to retain\n"); L.append("- "+a['evidence']+"\n")
    L.append("---\n")
    L.append(f"*{C.TITLE} · {C.COURSE_CODE} · Version {C.VERSION} · © 2026 {C.ORG}*")
    open(os.path.join(folder,"README.md"),"w").write("\n".join(L))

# ---- instruction DOCX -> PDF ---------------------------------------------
def _style(doc):
    st=doc.styles["Normal"]; st.font.name="Arial"; st.font.size=Pt(10.5)
def build_instructions(a,folder):
    doc=Document(); _style(doc)
    for s in doc.sections:
        s.top_margin=Inches(0.8); s.bottom_margin=Inches(0.8)
        s.left_margin=Inches(0.9); s.right_margin=Inches(0.9)
    logo=os.path.join(REPO,"courseware","assets","tertiary-infotech-logo.png")
    if os.path.exists(logo):
        p=doc.add_paragraph(); p.alignment=WD_ALIGN_PARAGRAPH.RIGHT
        p.add_run().add_picture(logo,height=Inches(0.42))
    h=doc.add_paragraph(); r=h.add_run(f"LAB {a['num']:02d} — INSTRUCTIONS")
    r.font.size=Pt(11); r.font.bold=True; r.font.color.rgb=RGBColor(0x1F,0x6F,0xEB)
    t=doc.add_paragraph(); r=t.add_run(a['title']); r.font.size=Pt(20); r.font.bold=True
    m=doc.add_paragraph(); r=m.add_run(
        f"Topic {a['topic']:02d}  ·  Day {a['day']}  ·  {a['level']}  ·  {a['apps']}  ·  Alignment {a['km']} · {a['am']}")
    r.font.size=Pt(9); r.font.color.rgb=RGBColor(0x5B,0x63,0x72)

    def head(txt):
        p=doc.add_paragraph(); r=p.add_run(txt); r.font.size=Pt(12); r.font.bold=True
        r.font.color.rgb=RGBColor(0x16,0x1B,0x26); p.paragraph_format.space_before=Pt(12)
        p.paragraph_format.space_after=Pt(4)
    head("Goal"); doc.add_paragraph(a['desc'])
    head("What you will build"); doc.add_paragraph(a['build'])
    head("Files supplied in this folder")
    for fn,d in a['files']: doc.add_paragraph(f"{fn} — {d}",style="List Bullet")
    head("Primary prompt")
    p=doc.add_paragraph(); r=p.add_run(a['prompt']); r.font.italic=True
    p.paragraph_format.left_indent=Inches(0.3)
    head("Detailed procedure")
    for i,(ttl,d) in enumerate(a['steps'],1):
        p=doc.add_paragraph(); r=p.add_run(f"Step {i}. {ttl}")
        r.font.bold=True; r.font.size=Pt(10.5); p.paragraph_format.space_before=Pt(8)
        p.paragraph_format.space_after=Pt(2)
        doc.add_paragraph(d)
    head("Acceptance test"); doc.add_paragraph(a['test'])
    head("Evidence to retain"); doc.add_paragraph(a['evidence'])
    f=doc.add_paragraph(); r=f.add_run(
        f"{C.TITLE} · {C.COURSE_CODE} · Version {C.VERSION} · © 2026 {C.ORG}")
    r.font.size=Pt(8); r.font.color.rgb=RGBColor(0x8A,0x93,0xA3)
    f.paragraph_format.space_before=Pt(18)
    dp=os.path.join(folder,f"LAB-{a['num']:02d}-Instructions.docx")
    doc.save(dp); return dp

def main():
    os.makedirs(LABS,exist_ok=True)
    docxs=[]
    for a in ACT:
        folder=os.path.join(LABS,f"lab-{a['num']:02d}-{slug(a['title'])}")
        os.makedirs(folder,exist_ok=True)
        os.makedirs(os.path.join(folder,"completed"),exist_ok=True)
        wbname=[f for f,_ in a['files'] if f.endswith(".xlsx")]
        wbname=wbname[0] if wbname else f"Lumina-Living-Lab-{a['num']:02d}.xlsx"
        build_workbook(a,os.path.join(folder,wbname))
        build_workbook(a,os.path.join(folder,"completed",wbname.replace(".xlsx","-completed.xlsx")))
        build_readme(a,folder)
        docxs.append(build_instructions(a,folder))
        print(f"lab {a['num']:02d} -> {os.path.basename(folder)}")
    # batch-convert instruction DOCX -> PDF
    soffice=shutil.which("soffice") or "/Applications/LibreOffice.app/Contents/MacOS/soffice"
    for d in docxs:
        subprocess.run([soffice,"--headless","--convert-to","pdf","--outdir",os.path.dirname(d),d],
                       check=False,capture_output=True)
    for d in docxs:
        if os.path.exists(d.replace(".docx",".pdf")): os.remove(d)
    print("done")

if __name__=="__main__": main()
